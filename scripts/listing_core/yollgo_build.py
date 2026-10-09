"""校验人工分组，生成可继续定价的 A1 候选表和变体映射。"""
import math
import hashlib
import csv
import io
import re
import unicodedata
from collections import Counter

from .common import CANDIDATE_FIELDS, CSV_FIELDS, Problem, atomic_bytes, component, inside, read_csv, read_json, require, write_csv, write_json
from .pricing import ean_ok
from .workspace import template_path
from .yollgo import number, other_candidates, selected_offer

MEASURES = ('weight_kg', 'len_cm', 'wid_cm', 'hei_cm')


def ean13(prefix):
    require(isinstance(prefix, str) and len(prefix) == 12 and prefix.isascii() and prefix.isdigit(),
            f'无法从 {prefix!r} 生成 13 位条码；请核对原条码和组合装件数。')
    total = sum(int(value) * (1 if index % 2 == 0 else 3) for index, value in enumerate(prefix))
    return prefix + str((-total) % 10)


def slug(value):
    text = unicodedata.normalize('NFKD', value).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', '-', text).strip('-')


def measurements(item, label):
    require(isinstance(item, dict), f'{label} 应为对象；请按 groups.json 示例填写。')
    values = {}
    for key in MEASURES:
        require(type(item.get(key)) in (int, float) and math.isfinite(item[key]) and item[key] > 0,
                f'{label}.{key} 必填且应是大于 0 的数字（重量 kg、尺寸 cm）；请补齐。')
        values[key] = item[key]
    require(type(item.get('estimado')) is bool, f'{label}.estimado 必须填写 true（估计）或 false（实测）。')
    return values


def text(obj, key, label, default=''):
    value = obj.get(key, default)
    require(isinstance(value, str), f'{label}.{key} 应为文字；请修正。')
    return value


def validate_groups(data, fetched):
    require(isinstance(data, dict) and data.get('version') == 1 and type(data.get('version')) is int,
            'groups.json.version 必须是 1；请按 references/数据格式.md 示例填写。')
    groups = data.get('groups')
    require(isinstance(groups, list) and groups, 'groups.json.groups 必须是非空数组；请至少填写一个组。')
    known, errors, plans = set(), [], []
    barcode_groups = {}
    for position, group in enumerate(groups, 1):
        label = f'groups.json 第 {position} 组'
        try:
            require(isinstance(group, dict), f'{label} 应为对象。')
            gid = group.get('group')
            require(isinstance(gid, str) and re.fullmatch(r'G[0-9]{2,}', gid), f'{label}.group 应为 G01、G02 等编号。')
            label += f'（{gid}）'
            require(gid not in known, f'{label}.group 重复；批次内组编号必须唯一。')
            known.add(gid)
            allowed = {'group', 'var1_name', 'var2_name', 'items', 'packs', 'notes'}
            require(not set(group) - allowed, f'{label} 有未知字段 {sorted(set(group) - allowed)}；请核对拼写。')
            names = [text(group, f'var{i}_name', label).strip() for i in (1, 2)]
            require(not names[1] or names[0], f'{label}.var2_name 需要同时填写 var1_name。')
            require(not names[1] or names[1] != names[0], f'{label} 两个变体名称不能相同。')
            notes = text(group, 'notes', label)
            packs = group.get('packs', [])
            require(isinstance(packs, list) and len(packs) <= 1 and all(type(qty) is int and 2 <= qty <= 9 for qty in packs),
                    f'{label}.packs 应为 [] 或一个 2～9 的整数，例如 [2]；每组一种组合装，生成 {gid}B。')
            items = group.get('items')
            require(isinstance(items, list) and items, f'{label}.items 必须是非空数组。')
            options, folders, entries = set(), set(), []
            for index, item in enumerate(items, 1):
                where = f'{label} 第 {index} 项'
                try:
                    values = measurements(item, where)
                    allowed_item = {'barcode', 'var1_value', 'var2_value', 'variante', 'notes', 'estimado', 'pack_measurements', *MEASURES}
                    require(not set(item) - allowed_item, f'{where} 有未知字段 {sorted(set(item) - allowed_item)}；请核对拼写。')
                    barcode = item.get('barcode')
                    require(isinstance(barcode, str) and len(barcode) == 13 and ean_ok(barcode),
                            f'{where}.barcode 必须是校验正确的 13 位条码字符串；请核对原条码，不要填写数字类型。')
                    require(barcode in fetched, f'{where}.barcode {barcode} 没有抓取记录；请先把它加入 barcodes.txt 并 fetch。')
                    offer = selected_offer(fetched[barcode])
                    require(barcode not in barcode_groups or barcode_groups[barcode] == gid,
                            f'{where}.barcode {barcode} 已分配给 {barcode_groups.get(barcode)}；同货号颜色请放在同一组。')
                    barcode_groups[barcode] = gid
                    variant = tuple(text(item, f'var{i}_value', where).strip() for i in (1, 2))
                    for i, (name, value) in enumerate(zip(names, variant), 1):
                        require(bool(name) == bool(value), f'{where}.var{i}_value 与组的 var{i}_name 必须同时有值或同时留空。')
                    require(variant not in options, f'{where} 变体取值 {variant} 重复；一个组合只能对应一个 SKU。')
                    options.add(variant)
                    manual = item.get('variante')
                    if manual is not None:
                        require(isinstance(manual, str) and re.fullmatch(r'[a-z0-9]+(?:-[a-z0-9]+)*', manual),
                                f'{where}.variante 只能用小写英文字母、数字和短横线，不要空格或重音。')
                        folder = component(manual, f'{where}.variante')
                        require(folder not in folders, f'{where}.variante {folder} 重复；请换一个文件夹名。')
                    else:
                        base = slug('-'.join(variant)) or 'unico'
                        folder, suffix = base, 2
                        while folder in folders:
                            folder, suffix = f'{base}-{suffix}', suffix + 1
                        component(folder, f'{where}.variante')
                    folders.add(folder)
                    overrides = item.get('pack_measurements', {})
                    require(isinstance(overrides, dict), f'{where}.pack_measurements 应按件数字符串填写对象。')
                    require(set(overrides) <= {str(q) for q in packs}, f'{where}.pack_measurements 只能填写本组 packs 中的件数。')
                    for qty, override in overrides.items():
                        measurements(override, f'{where}.pack_measurements.{qty}')
                        require(not set(override) - {*MEASURES, 'estimado'}, f'{where}.pack_measurements.{qty} 有未知字段；请核对拼写。')
                    entries.append(dict(item=item, barcode=barcode, offer=offer, values=values, variant=variant,
                                        folder=folder, note=text(item, 'notes', where), label=where))
                except Problem as exc:
                    errors.extend(error if error.startswith(where) else f'{where}：{error}' for error in exc.errors)
            plans.append(dict(group=gid, names=names, notes=notes, packs=packs, entries=entries))
        except Problem as exc:
            errors.extend(exc.errors)
    if errors:
        raise Problem(errors)
    return plans


def reserved_eans(ws, batch, shop):
    reserved = {}
    for batch_name, row in other_candidates(ws, batch):
        if row['ean']:
            reserved.setdefault(row['ean'], f'批次 {batch_name} / 组 {row["group"]}')
    import openpyxl
    from .headers import existing_eans
    # Missing template would leave an entire collision source unchecked: stop.
    # Official templates can declare an undersized dimension. Normal mode reads
    # the actual XML cells, whereas read_only would miss EANs outside that range.
    wb = openpyxl.load_workbook(template_path(ws, shop), keep_vba=True)
    try:
        for code in existing_eans(wb):
            reserved.setdefault(code, '模板已存在条码清单')
        require('Pantilla' in wb.sheetnames, '模板缺少 Pantilla；请重新下载官方模板。')
        sheet = wb['Pantilla']
        columns = [c.column for c in next(sheet.iter_rows(min_row=1, max_row=1)) if c.value == 'Código EAN']
        require(len(columns) == 1, '模板 Código EAN 表头缺失或重复；请换完整官方模板。')
        for cells in sheet.iter_rows(min_row=5, min_col=columns[0], max_col=columns[0]):
            value = cells[0].value
            if value is not None:
                code = str(int(value)) if type(value) in (int, float) and float(value).is_integer() else str(value).strip()
                reserved.setdefault(code, '模板 Pantilla 已有商品')
    finally:
        wb.close()
    return reserved


def build(ws, batch, shop, force=False):
    target = inside(batch, batch / 'candidates.csv')
    require(force or not target.exists() or not read_csv(target),
            'candidates.csv 已有商品，不会覆盖；确认需要重新生成时加 --force，然后重新运行 price。')
    for name in ('variantes.csv', 'build.json'):
        inside(batch, batch / name)
    raw = read_json(batch / 'yollgo.json')
    require(isinstance(raw, dict) and raw.get('version') == 1 and isinstance(raw.get('items'), list),
            'yollgo.json 格式不对；请先运行 fetch。')
    fetched = {}
    for item in raw['items']:
        require(isinstance(item, dict) and isinstance(item.get('barcode'), str) and item['barcode'] not in fetched,
                'yollgo.json.items 条码缺失或重复；请重新 fetch。')
        fetched[item['barcode']] = item
    plans = validate_groups(read_json(batch / 'groups.json'), fetched)
    reserved = reserved_eans(ws, batch, shop)
    # Reserve ALL source EANs before any color is generated. A later item's
    # original barcode must never get stolen by an earlier synthetic color.
    originals = {entry['barcode'] for plan in plans for entry in plan['entries']}
    for code in sorted(originals):
        require(code not in reserved, f'原条码 {code} 撞号，已在 {reserved.get(code)}；请核对分组或已上架商品，不会自动改原条码。')
    used = dict(reserved)
    used.update({code: '本批次原条码' for code in originals})
    planned_packs = {ean13('2' + str(qty) + entry['barcode'][-10:])
                     for plan in plans for qty in plan['packs'] for entry in plan['entries']}
    rows, mappings, warnings, first = [], [], [], set()
    for plan in plans:
        gid, single_rows = plan['group'], []
        for entry in plan['entries']:
            code = entry['barcode']
            if code in first:
                options = [ean13(code[:11] + str(index)) for index in range(1, 10)]
                ean = next((option for option in options if option not in used and option not in planned_packs), None)
                require(ean is not None, f'{entry["label"]} 条码 {code} 的 1～9 颜色号都已占用；请减少颜色或核对已有条码。')
            else:
                ean = code
                first.add(code)
            used[ean] = f'{gid}/{entry["folder"]}'
            offer, item = entry['offer'], entry['item']
            product, sid = offer['product'], offer['shop_id']
            supplier = shop['suppliers'].get(sid, {})
            price = number(product.get('precio'), f'{entry["label"]} precio', positive=True)
            unit = price / 1.21 if supplier.get('price_includes_iva', False) else price
            minimum = number(product.get('baozhuangshu', 1), f'{entry["label"]} baozhuangshu', positive=True)
            notes = [f'起订包装数：{minimum:g}', fetched[code].get('note', ''), plan['notes'], entry['note'],
                     '重量和包装尺寸：估计' if item['estimado'] else '重量和包装尺寸：实测']
            if sid not in shop['suppliers']:
                warnings.append(f'{gid} / 条码 {code} 的批发商 {sid} 未登记在 shop.json.suppliers，制造商会空；fill 前请补齐。')
            row = dict(group=gid, shop=sid, art_id=str(product['artId']), ean=ean, unit_ean=code,
                       src_name=product.get('namecn') or product.get('namees', ''),
                       var1_name=plan['names'][0], var1_value=entry['variant'][0],
                       var2_name=plan['names'][1], var2_value=entry['variant'][1], pack_qty=1,
                       unit_cost_ex_iva=unit, img_hash=product.get('imageHash', ''), **entry['values'],
                       market_low='', market_high='', notes='；'.join(str(note) for note in notes if note))
            single_rows.append((row, entry))
            rows.append(row)
            mappings.append(dict(group=gid, art_id=row['art_id'], ean=ean, variante=entry['folder']))
        for qty in plan['packs']:
            pack_group = gid + 'B'
            for original, entry in single_rows:
                code = ean13('2' + str(qty) + original['ean'][-10:])
                require(code not in used, f'{entry["label"]} / {pack_group} 的 {qty} 件装条码 {code} 撞号，已在 {used.get(code)}；请更改组合装件数或核对已有 SKU。')
                used[code] = f'{pack_group}/{entry["folder"]}'
                override = entry['item'].get('pack_measurements', {}).get(str(qty))
                values = measurements(override, f'{entry["label"]} 组合装') if override else {
                    **entry['values'], 'weight_kg': entry['values']['weight_kg'] * qty,
                    'hei_cm': entry['values']['hei_cm'] * qty}
                row = dict(original, group=pack_group, ean=code, pack_qty=qty,
                           unit_cost_ex_iva=original['unit_cost_ex_iva'] * qty, **values)
                row['notes'] += f'；{qty} 件装；组合装重量和包装尺寸：' + ('估计（默认叠放）' if override is None else ('估计' if override['estimado'] else '实测'))
                rows.append(row)
                mappings.append(dict(group=pack_group, art_id=row['art_id'], ean=code, variante=entry['folder']))
    require(rows, 'groups.json 没有可生成的 SKU；请补齐分组。')
    omitted = sorted(set(fetched) - originals)
    if omitted:
        warnings.append('以下已抓取条码未放入本次分组：' + '、'.join(omitted))
    # Cost is the WHOLE SKU/pack per A2. Explicit provenance lets price retain
    # A1's per-unit multiplication for legacy batches without guessing by name.
    stream = io.StringIO(newline='')
    writer = csv.DictWriter(stream, fieldnames=CANDIDATE_FIELDS)
    writer.writeheader()
    writer.writerows(rows)
    payload = stream.getvalue().encode('utf-8-sig')
    # Install provenance before pack costs. On a first build interrupted while
    # writing files, price must see a hash mismatch instead of legacy unit math.
    write_json(batch / 'build.json', dict(version=1, cost_basis='pack',
               candidates_sha256=hashlib.sha256(payload).hexdigest()))
    write_csv(batch / 'variantes.csv', mappings, CSV_FIELDS['variantes.csv'])
    atomic_bytes(target, payload)
    return dict(count=len(rows), groups=len(Counter(row['group'] for row in rows)), output=str(target), warnings=warnings,
                message='候选商品和变体映射已生成；重量尺寸估计写在备注里，请继续运行 price。')
