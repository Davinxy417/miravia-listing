"""检查与填表共用准备流程；只有验证通过后才原子替换输出。"""
import tempfile
from collections import defaultdict
from pathlib import Path
from zipfile import ZipFile

import openpyxl

from .common import Problem, inside, read_csv, read_json, require
from .content import content_errors, number, resolve_content, seller_sku
from .headers import HEADERS, dropdown, existing_eans, resolve_columns
from .images import group_images, read_image_urls, read_variantes, variant_key
from .pricing import ean_ok
from .xlsm import patch_sheet, sheet_part, verify_integrity



DESCRIPTION_MAX = 3000  # 模板 Descripción 上限


def description_html(text, gallery):
    """图床图插进描述:02 卖点图放最前,03 细节和 04/07/08 场景图放文字后;友购占位图不插。超长就从后往前少插。"""
    real = lambda i: i < len(gallery) and gallery[i] and 'freex.es' not in gallery[i]
    top = [gallery[1]] if real(1) else []
    bottom = [gallery[i] for i in (2, 3, 6, 7) if real(i)]
    tag = lambda url: f'<p><img src="{url}" style="width:100%"/></p>'
    while True:
        html = ''.join(map(tag, top)) + text + ''.join(map(tag, bottom))
        if len(html) <= DESCRIPTION_MAX or not (top or bottom): return html
        (bottom or top).pop()

def prepare(template, batch, shop):
    errors, warnings = [], set()
    def read(fn, fallback):
        try: return fn()
        except (ValueError, KeyError, OSError) as exc:
            errors.extend(exc.errors if isinstance(exc, Problem) else [f'{exc}；请补齐该文件或字段后重试。'])
            return fallback
    rows = read(lambda: read_csv(batch / 'priced.csv'), [])
    candidates = read(lambda: read_csv(batch / 'candidates.csv'), [])
    require(rows, 'priced.csv 没有可填商品；请先运行 price。')
    raw_content = read(lambda: read_json(batch / 'content.json'), {})
    content, issues = resolve_content(raw_content, [r['group'] for r in rows])
    errors.extend(issues)
    mappings = read(lambda: read_variantes(batch / 'variantes.csv'), {})
    images = read(lambda: read_image_urls(batch / 'images.csv', mappings), {})
    if len(rows) != len(candidates) or any(any(r.get(k) != v for k, v in c.items()) for c, r in zip(candidates, rows)):
        errors.append('priced.csv 与 candidates.csv 不一致；候选商品改动后请重新运行 price。')
    if (batch / 'candidates.csv').exists() and (batch / 'priced.csv').stat().st_mtime_ns < (batch / 'candidates.csv').stat().st_mtime_ns:
        warnings.add('priced.csv 比 candidates.csv 旧；请确认是否需要重新定价。')
    wb = openpyxl.load_workbook(template, keep_vba=True)
    try:
        columns = resolve_columns(wb)
        lists = {field: read(lambda f=field: dropdown(wb, f, columns), set()) for field in ('category', 'warning', 'manufacturer', 'eu', 'hazard')}
        blocked = read(lambda: existing_eans(wb), set())
        if 'Ninguno' not in lists['hazard']:
            errors.append('模板危险品下拉没有 Ninguno；请下载适用的官方模板。')
        by_group, seen_eans, seen_skus = defaultdict(list), set(), set()
        for group, item in content.items():
            errors.extend(content_errors(group, item))
            for field in ('category', 'warning'):
                if item.get(field) not in lists[field]:
                    errors.append(f'{group} 的 {field} 不在模板下拉中；请使用完整原文，可用 categories 搜索类目。')
            if group.endswith('B') and isinstance(raw_content.get(group[:-1]), dict):
                base = raw_content[group[:-1]]
                if any(item.get(k) != base.get(k) for k in ('warning', 'warning_text')):
                    errors.append(f'{group} 安全提示与主组不同；请核对并保持一致。')
        for n, row in enumerate(rows, 2):
            group, ean = row['group'], row['ean']
            label = f'priced.csv 第 {n} 行 {group}/{row["art_id"]}'
            by_group[group].append(row)
            if not ean_ok(ean): errors.append(f'{label} 条码 {ean!r} 无效；请核对数字、长度和校验位。')
            if ean in seen_eans or ean in blocked: errors.append(f'{label} 条码 {ean} 重复或已在店铺使用；请换用正确且唯一的条码。')
            seen_eans.add(ean)
            if not row['img_hash']: errors.append(f'{label} 缺 img_hash；请补齐原始商品图片标识。')
            if variant_key(row) not in mappings: errors.append(f'{label} 缺变体目录映射；请补齐 variantes.csv。')
            maker = shop['suppliers'].get(row['shop'])
            if not maker:
                errors.append(f'{label} 商家 {row["shop"]} 未配置；请在 shop.json.suppliers 中补齐公司和下拉原文。')
            else:
                for key, field in (('fabricante', 'manufacturer'), ('responsable', 'eu')):
                    if maker[key] not in lists[field]:
                        errors.append(f'商家 {row["shop"]} 的 {key} 不在模板下拉中；请更新配置原文或重新下载已审核模板。')
            numeric = {}
            for field in ('pack_qty', 'price', 'original_price', 'min_safe_price', 'weight_kg', 'len_cm', 'wid_cm', 'hei_cm'):
                try:
                    numeric[field] = number(row[field], f'{label}.{field}', integer=field in ('pack_qty', 'len_cm', 'wid_cm', 'hei_cm'), money=field in ('price', 'original_price'))
                except ValueError as exc: errors.append(str(exc))
            if all(k in numeric for k in ('price', 'min_safe_price')) and numeric['price'] < numeric['min_safe_price']:
                errors.append(f'{label} 售价低于保本价；请重新定价或核对成本。')
            if all(k in numeric for k in ('price', 'original_price')) and numeric['price'] != numeric['original_price']:
                errors.append(f'{label} 原价应等于售价（本流程不填折扣价）；请把 original_markup 设为 1 后重新定价。')
            if group in content and 'pack_qty' in numeric:
                sku = seller_sku(row, shop, content[group])
                if sku in seen_skus or len(sku) > 200: errors.append(f'{label} SKU 重复或过长；请检查 art_id、组合件数和 sku_include_ean。')
                seen_skus.add(sku)
        output = []
        for group, variants in sorted(by_group.items()):
            if len(variants) > 1:
                names = {(r['var1_name'], r['var2_name']) for r in variants}
                if len(names) != 1 or any(not a or a == b for a, b in names):
                    errors.append(f'{group} 变体名称缺失、不一致或重复；请统一 var1_name/var2_name。')
                combos = set()
                for row in variants:
                    combo = row['var1_value'], row['var2_value']
                    if not combo[0] or bool(row['var2_name']) != bool(combo[1]) or combo in combos:
                        errors.append(f'{group}/{row["ean"]} 变体选项缺失或重复；请补齐且确保组合唯一。')
                    combos.add(combo)
            if all(variant_key(r) in mappings for r in variants):
                read(lambda: group_images(group, variants, mappings, images, batch / 'images', warnings), None)
        if len(rows) > 32764: errors.append('商品超过模板允许的 32764 行；请拆成多个批次。')
        if errors: raise Problem(list(dict.fromkeys(errors)), sorted(warnings))
        for group, variants in sorted(by_group.items()):
            item = content[group]
            gallery, principals = group_images(group, variants, mappings, images, batch / 'images', warnings)
            for row in variants:
                maker = shop['suppliers'][row['shop']]
                fields = dict(group=group, category=item['category'], title=item['title'], brand=shop['brand'],
                              attributes=item['attributes'], description=description_html(item['description'], gallery), warning=item['warning'],
                              warning_text=item['warning_text'], variant_image=principals[mappings[variant_key(row)]],
                              ean=row['ean'], sku=seller_sku(row, shop, item), price=number(row['price'], 'price'),
                              original_price=number(row['price'], 'price'), stock=shop['stock_default'],
                              manufacturer=maker['fabricante'], eu=maker['responsable'], hazard='Ninguno')
                fields.update({f'image{i}': url for i, url in enumerate(gallery, 1)})
                if len(variants) > 1:
                    fields.update({key: row[key] for key in ('var1_name', 'var1_value', 'var2_name', 'var2_value')})
                for key in ('weight_kg', 'len_cm', 'wid_cm', 'hei_cm'):
                    fields[key] = number(row[key], key, integer=key != 'weight_kg')
                values = [None] * wb['Pantilla'].max_column
                for key, value in fields.items(): values[columns[key] - 1] = value
                output.append(values)
        return output, sorted(warnings), columns, wb['Pantilla'].max_column
    finally:
        wb.close()


def inspect_local_images(batch):
    """只读提示；本地文件并不是已托管 URL 的必要条件。"""
    warnings = []
    manifest = batch / 'image_files.csv'
    if not manifest.exists(): return ['缺 image_files.csv；有本地图片后请运行 images-collect。']
    for row in read_csv(manifest):
        raw = row['local_path'].replace('\\', '/')
        path = Path(raw)
        require(raw and not path.is_absolute() and '..' not in path.parts and ':' not in raw,
                f'image_files.csv 路径 {raw!r} 不正确；请用 images/ 下的相对路径。')
        if not (batch / 'images' / path).is_file(): warnings.append(f'本地图片不存在：{raw}；请检查图片目录或重新收集。')
    return warnings


def run(template, batch, shop, check=False):
    local_warnings, local_errors = [], []
    if check:
        try: local_warnings = inspect_local_images(batch)
        except (ValueError, OSError) as exc:
            local_errors = exc.errors if isinstance(exc, Problem) else [f'图片清单检查失败：{exc}；请检查 image_files.csv。']
    try:
        values, warnings, columns, width = prepare(template, batch, shop)
    except Problem as exc:
        raise Problem(exc.errors + local_errors, exc.warnings + local_warnings) from exc
    if check:
        warnings.extend(local_warnings)
        if local_errors: raise Problem(local_errors, warnings)
        return dict(count=len(values), warnings=warnings, upload_ready=not warnings, message='检查完成，未写入文件。')
    out = inside(batch, batch / 'output/miravia_upload.xlsm')
    write_upload(template, out, values, columns, width)
    return dict(count=len(values), groups=len({r[columns['group']-1] for r in values}), output=str(out), warnings=warnings, upload_ready=not warnings)


def write_upload(template, out, values, columns, width):
    """完整或分批表共用同一 ZIP 复制及完整性检查。"""
    out = Path(out)
    require(out != Path(template).resolve(), '输出不能覆盖模板；请改用独立 output/ 目录。')
    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=out.parent, suffix='.xlsm', delete=False) as f:
        pending = Path(f.name)
    try:
        with ZipFile(template) as source, ZipFile(pending, 'w') as result:
            part = sheet_part(source)
            edited = patch_sheet(source, part, values, columns, width)
            result.comment = source.comment
            for member in source.infolist():
                result.writestr(member, edited if member.filename == part else source.read(member.filename))
        verify_integrity(template, pending)
        pending.replace(out)
    finally:
        pending.unlink(missing_ok=True)
