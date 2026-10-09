"""唯一的字段→模板表头对照；隐藏数据同样按字段标识查找。"""
from openpyxl.utils import get_column_letter
from .common import Problem, require

HEADERS = {
    'group': 'Group No', 'category': 'Categoría', 'title': 'Nombre del producto',
    **{f'image{i}': f'Imágenes de producto{i}' for i in range(1, 9)},
    'brand': 'Marca', 'attributes': 'Atributos adicionales', 'description': 'Descripción',
    'warning': '¿El producto cuenta con advertencia de seguridad?',
    'warning_text': 'Contenido de la advertencia de seguridad',
    'var1_name': 'Variation Name1', 'var1_value': 'Option for Variation1',
    'variant_image': 'Image per Variation',
    'var2_name': 'Variation Name2', 'var2_value': 'Option for Variation2',
    'ean': 'Código EAN', 'sku': 'SKU de vendedor', 'original_price': 'Precio original',
    'price': 'Precio en España', 'discount_price': 'Precio con descuento en España',
    'stock': 'Stock', 'manufacturer': 'Fabricante', 'eu': 'Persona Responsable de la UE',
    'weight_kg': 'Peso del paquete', 'len_cm': 'Longitud del paquete',
    'wid_cm': 'Ancho del paquete', 'hei_cm': 'Altura del paquete', 'hazard': 'Materiales peligrosos',
}
HIDDEN_FIELDS = {
    'category': 'categoryPath', 'warning': 'safetyWarningRequirement',
    'manufacturer': 'sku.Manufacturer_info', 'eu': 'sku.EU_Responsible',
    'hazard': 'dangerousGoods', 'ean': 'sku.ean_code',
}


def resolve_columns(wb):
    require('Pantilla' in wb.sheetnames, '模板缺少 Pantilla 工作表；请重新下载官方 .xlsm 模板。')
    sheet, columns, errors = wb['Pantilla'], {}, []
    for field, header in HEADERS.items():
        matches = []
        for cell in sheet[1]:
            text = str(cell.value or '').strip()
            if text == header or (field == 'stock' and text.startswith(header + '<')):
                matches.append(cell.column)
        if len(matches) != 1:
            errors.append(f'模板表头 {header!r} 找到 {len(matches)} 列 {matches}；请确认使用完整官方模板，且表头不重名。')
        else:
            columns[field] = matches[0]
    if errors: raise Problem(errors)
    col = columns['variant_image']
    instructions = ' '.join(str(sheet.cell(r, col).value or '') for r in (3, 4)).lower()
    require('fondo blanco' in instructions and '1:1' in instructions,
            f'模板 {get_column_letter(col)} 列变体图说明已变化；请核对是否仍要求白底 1:1 图片。')
    return columns


def hidden_column(wb, field, sheet_name, header_row):
    require(sheet_name in wb.sheetnames, f'模板缺少隐藏表 {sheet_name}；请重新下载官方模板。')
    matches = [c.column for c in wb[sheet_name][header_row] if c.value == HIDDEN_FIELDS[field]]
    require(len(matches) == 1, f'模板 {sheet_name} 的字段 {HIDDEN_FIELDS[field]} 缺失或重复；请换用完整官方模板。')
    return matches[0]


def dropdown(wb, field, columns=None):
    # Visible positions and hidden list positions need not move together.
    if columns is None: columns = resolve_columns(wb)
    require(field in columns, f'模板未找到 {HEADERS[field]} 表头；请检查模板。')
    col = hidden_column(wb, field, 'Pantilla_hide', 3)
    values = set()
    for cells in wb['Pantilla_hide'].iter_rows(min_row=7, min_col=col, max_col=col):
        value = cells[0].value
        if value == '####splitter line####': break
        if value is not None: values.add(str(value))
    require(bool(values), f'模板 {HEADERS[field]} 下拉为空；请在后台登记资料并重新下载模板。')
    return values


def existing_eans(wb):
    col = hidden_column(wb, 'ean', 'global_validation_hide', 1)
    return {str(cells[0].value).strip() for cells in wb['global_validation_hide'].iter_rows(min_row=3, min_col=col, max_col=col) if cells[0].value is not None}
