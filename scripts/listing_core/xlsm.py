"""复制 ZIP 成员，仅替换 Pantilla 的 sheetData 和 dimension。"""
import posixpath
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile
from xml.sax.saxutils import escape
import openpyxl
from openpyxl.utils import get_column_letter
from .common import require
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
def sheet_part(archive, name="Pantilla"):
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    sheet = next(s for s in workbook.find("m:sheets", NS) if s.get("name") == name)
    target = next(r.get("Target") for r in relationships if r.get("Id") == sheet.get(f"{{{REL_NS}}}id"))
    return posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)

DATA_RE = re.compile(r"(<sheetData\b[^>]*>)(.*?)(</sheetData>)", re.S)

DIM_RE = re.compile(r'<dimension\b[^>]*/>')

ROW_RE = re.compile(r'<row\b[^>]*\br="(\d+)"[^>]*(?:/>|>.*?</row>)', re.S)

def sheet_without_data(xml):
    return DIM_RE.sub('<dimension/>', DATA_RE.sub(r'\1\3', xml))

def patch_sheet(archive, part, rows, columns, width):
    xml = archive.read(part).decode("utf-8")
    require(DATA_RE.search(xml) is not None, "模板数据区格式不支持；请使用原始官方 .xlsm 模板。")
    root = ET.fromstring(xml)
    xfs = ET.fromstring(archive.read("xl/styles.xml")).find("m:cellXfs", NS)
    text_style = next((i for i, xf in enumerate(xfs) if xf.get("numFmtId") == "49"), None)
    require(text_style is not None, "模板没有条码文本样式（@）；请重新下载官方模板。")
    numeric_style = next((i for i, xf in enumerate(xfs) if xf.get("numFmtId", "0") == "0"), 0)
    money_style = next((i for i, xf in enumerate(xfs) if xf.get("numFmtId") == "2"), numeric_style)
    column_styles = {}
    for col in root.findall("m:cols/m:col", NS):
        for idx in range(int(col.get("min")), int(col.get("max")) + 1):
            column_styles[idx] = int(col.get("style", "0"))
    old_data = DATA_RE.search(xml).group(2)
    # Keep the four original instruction rows exactly as serialized.
    header_rows = [m.group(0) for m in ROW_RE.finditer(old_data) if int(m.group(1)) < 5]
    require(len(header_rows) == 4, "模板应保留原始 1～4 行表头及说明；请恢复官方模板。")
    new_rows = []
    for row_num, values in enumerate(rows, 5):
        cells = []
        for col, value in enumerate(values, 1):
            if value is None or value == "":
                continue
            address = f"{get_column_letter(col)}{row_num}"
            if isinstance(value, str):
                require(not re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value), f"{address} 含不支持的控制字符；请清理对应文案。")
                style = text_style if col == columns["ean"] else column_styles.get(col, text_style)
                cells.append(f'<c r="{address}" s="{style}" t="inlineStr"><is><t xml:space="preserve">{escape(value)}</t></is></c>')
            else:
                # Keep the template's unlocked stock/measurement styles: the
                # workbook is protected, so style 0 would lock numeric cells.
                style = money_style if col in {columns[k] for k in ("original_price", "price", "discount_price")} else column_styles.get(col, numeric_style)
                cells.append(f'<c r="{address}" s="{style}" t="n"><v>{value}</v></c>')
        new_rows.append(f'<row r="{row_num}">' + "".join(cells) + '</row>')
    body = "\n" + "\n".join(header_rows + new_rows) + "\n"
    xml = DATA_RE.sub(lambda m: m.group(1) + body + m.group(3), xml, count=1)
    xml, count = DIM_RE.subn(f'<dimension ref="A1:{get_column_letter(width)}{len(rows) + 4}"/>', xml, count=1)
    require(count == 1, "模板缺少有效的数据范围声明；请重新下载官方模板。")
    ET.fromstring(xml)
    return xml.encode("utf-8")

def verify_integrity(template, output):
    with ZipFile(template) as source, ZipFile(output) as result:
        require(result.testzip() is None, "生成文件压缩包校验失败；原输出未覆盖，请重试或检查磁盘。")
        require(len(result.namelist()) == len(set(result.namelist())), "生成文件存在重复压缩成员；请检查模板，原输出未覆盖。")
        require(set(source.namelist()) == set(result.namelist()), "模板内部文件清单改变；请检查填表程序，原输出未覆盖。")
        require("xl/vbaProject.bin" in result.namelist(), "文件缺少宏；请使用带宏的官方 .xlsm 模板。")
        part = sheet_part(source)
        for name in source.namelist():
            a, b = source.read(name), result.read(name)
            if name == part:
                require(sheet_without_data(a.decode()) == sheet_without_data(b.decode()),
                        "Pantilla 非数据区发生变化；请检查填表程序，原输出未覆盖。")
                headers = lambda data: [m.group(0) for m in ROW_RE.finditer(DATA_RE.search(data.decode()).group(2)) if int(m.group(1)) < 5]
                require(headers(a) == headers(b), "模板说明行改变；请检查填表程序，原输出未覆盖。")
            else:
                require(a == b, f"模板成员改变：{name}；请检查填表程序，原输出未覆盖。")
            if name.endswith((".xml", ".rels")):
                ET.fromstring(b)
    before = openpyxl.load_workbook(template, keep_vba=True)
    after = openpyxl.load_workbook(output, keep_vba=True)
    try:
        require(before.sheetnames == after.sheetnames, "工作表清单发生变化；请检查填表程序。")
        for sheet in before:
            require(sheet.sheet_state == after[sheet.title].sheet_state, f"隐藏表状态改变：{sheet.title}；请检查填表程序。")
        a = before["Pantilla"].data_validations.dataValidation
        b = after["Pantilla"].data_validations.dataValidation
        require(len(a) == len(b) and len(a) > 0, "数据校验规则缺失或数量改变；请检查模板和填表程序。")
        require([ET.tostring(v.to_tree()) for v in a] == [ET.tostring(v.to_tree()) for v in b],
                "数据校验规则内容发生变化；请检查填表程序。")
    finally:
        before.close()
        after.close()
