"""A4 离线验收：只有 .tmp 内的 file:// bare 图床，gh/HEAD 全部替身。"""
import contextlib
import hashlib
import io
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path
from unittest.mock import patch

import openpyxl
from PIL import Image
from helpers import FIXTURES, ROOT, cli, csv_rows, json_write, workspace
import mlist
from listing_core import image_host as host
from listing_core.common import Problem, read_json, write_csv, CSV_FIELDS
from listing_core.image_layout import SLOTS, output_dir, variant_folders
from listing_core.workspace import load_shop
from smoke_test import independent_zip


def local_git(directory, *args):
    result = subprocess.run(['git', '-c', 'user.name=Offline Test', '-c', 'user.email=test@example.invalid', *args],
                            cwd=directory, capture_output=True, text=True, encoding='utf-8', errors='replace')
    assert result.returncode == 0, (args, result.stderr)
    return result.stdout.strip()


def captured(ws, *args):
    stream = io.StringIO()
    with contextlib.redirect_stdout(stream):
        code = mlist.main([*map(str, args), '--ws', str(ws), '--json'])
    result = json.loads(stream.getvalue())
    assert not result['errors'] or not result['ok'], result
    return code, result


def snap(directory):
    return {str(p.relative_to(directory)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in directory.rglob('*') if p.is_file()}


def jpeg(colour):
    stream = io.BytesIO()
    Image.new('RGB', (24, 24), colour).save(stream, 'JPEG')
    return stream.getvalue()


class FakeGitHub:
    def __init__(self, root):
        self.root, self.remotes, self.calls, self.checked = root, {}, [], []
        self.private = False

    def create(self, repo, seed=False):
        bare = self.root / (repo.replace('/', '-') + '.git')
        assert bare.resolve().is_relative_to((ROOT / '.tmp').resolve())
        local_git(self.root, 'init', '--bare', '--initial-branch=main', str(bare))
        self.remotes[repo] = bare.as_uri()
        if seed:
            clone = self.root / (repo.replace('/', '-') + '-seed')
            local_git(self.root, 'clone', bare.as_uri(), str(clone))
            (clone / 'README.md').write_text('保留原有 README\n', 'utf-8')
            (clone / 'old-batch').mkdir()
            (clone / 'old-batch/legacy.jpg').write_bytes(jpeg('red'))
            local_git(clone, 'add', '.')
            local_git(clone, 'commit', '-m', 'existing images')
            local_git(clone, 'push', 'origin', 'main')
        return bare

    def gh(self, *args, allow_failure=False):
        self.calls.append(args)
        if args[:2] == ('api', 'user'):
            return subprocess.CompletedProcess(args, 0, 'offline-owner\n', '')
        if args[:2] == ('repo', 'view'):
            repo = args[2]
            if repo not in self.remotes:
                return subprocess.CompletedProcess(args, 1, '', 'Could not resolve to a Repository')
            return subprocess.CompletedProcess(args, 0, json.dumps(dict(isPrivate=self.private, nameWithOwner=repo)), '')
        if args[:2] == ('repo', 'create'):
            assert '--public' in args
            self.create(args[2])
            return subprocess.CompletedProcess(args, 0, '', '')
        raise AssertionError(('unexpected gh', args))

    def remote_url(self, config):
        assert config['repo'] in self.remotes
        return self.remotes[config['repo']]

    def head(self, url):
        assert url.startswith(('https://raw.githubusercontent.com/offline-owner/', 'https://cdn.example.invalid/'))
        self.checked.append(url)
        return dict(url=url, ok=True, status=200, content_type='image/jpeg')


def make_ws():
    ws = workspace('publish-')
    cli(ws, 'init')
    shutil.copyfile(FIXTURES / 'shop.json', ws / 'shop.json')
    shop = read_json(ws / 'shop.json')
    shutil.copyfile(FIXTURES / 'template' / shop['template'], ws / 'template' / shop['template'])
    name = '2026-10-冬季第一批'
    cli(ws, 'new-batch', name)
    batch = ws / 'batches' / name
    shutil.copytree(FIXTURES / 'batch', batch, dirs_exist_ok=True)
    for group, variants in variant_folders(batch / 'variantes.csv').items():
        for variant in variants:
            directory = output_dir(batch / 'images', group, variant)
            directory.mkdir(parents=True, exist_ok=True)
            for filename in [*SLOTS.values(), 'variante.jpg']:
                (directory / filename).write_bytes(jpeg('blue'))
    cli(ws, 'price', '--batch', name)
    return ws, batch


def mark(batch, items, result='ok'):
    records = read_json(batch / 'image_review.json') if (batch / 'image_review.json').exists() else {'version': 1, 'reviews': {}}
    for item in items:
        records['reviews'][item['id']] = dict(result=result, sha256=item['sha256'], note='测试' if result=='redo' else '', at='2026-10-09T00:00:00+00:00')
    json_write(batch / 'image_review.json', records)


def main():
    ws, batch = make_ws()
    fake = FakeGitHub(ws)
    bare = fake.create('offline-owner/images', seed=True)
    with patch.object(host, 'gh', fake.gh), patch.object(host, 'remote_url', fake.remote_url), patch.object(host, 'head_image', fake.head):
        before = snap(ws)
        with patch.object(host, 'command', side_effect=AssertionError('预览不能调用任何进程')):
            code, result = captured(ws, 'image-host-setup', '--repo', 'images')
        assert code == 2 and result['needs_confirm'] and snap(ws) == before and not fake.calls
        code, result = captured(ws, 'image-host-setup', '--repo', 'offline-owner/images', '--yes')
        assert code == 0, result
        shop = load_shop(ws)
        clone = host.clone_path(ws, shop['image_host']['repo'])
        assert (clone / 'README.md').read_text('utf-8') == '保留原有 README\n'
        legacy = (clone / 'old-batch/legacy.jpg').read_bytes()
        before_head = local_git(clone, 'rev-parse', 'HEAD')
        code, _ = captured(ws, 'image-host-setup', '--repo', 'offline-owner/images', '--yes')
        assert code == 0 and local_git(clone, 'rev-parse', 'HEAD') == before_head
        # Missing repo is created public, empty clone gets the requested README.
        empty_ws = workspace('publish-empty-')
        cli(empty_ws, 'init')
        code, result = captured(empty_ws, 'image-host-setup', '--repo', 'new-images', '--yes')
        assert code == 0, result
        assert (empty_ws / '.image-host/new-images/README.md').read_text('utf-8').strip() == '本仓库只存放店铺商品图'
        assert any(a[:3] == ('repo','create','offline-owner/new-images') for a in fake.calls)
        fake.private = True
        code, result = captured(ws, 'image-host-setup', '--repo', 'offline-owner/images', '--yes')
        assert code == 1 and any('公开' in e for e in result['errors'])
        fake.private = False
        print('setup：未确认零写入/零进程，假 gh 新建公开仓库，已有 README/旧图保留，拒绝私有仓库通过')
        all_items, _ = host.selection(batch, True)
        mark(batch, all_items[:1])
        # Failure isolation is exercised by test_a6; here other images stay pending.
        before = snap(ws)
        with patch.object(host, 'command', side_effect=AssertionError('预览不能调用进程')), patch.object(host, 'head_image', side_effect=AssertionError('预览不能联网')):
            code, result = captured(ws, 'publish-images', '--batch', batch.name)
        assert code == 2 and result['needs_confirm'] and result['count'] == 1 and result['total_bytes'] == all_items[0]['size']
        assert result['skipped'] and snap(ws) == before
        code, result = captured(ws, 'publish-images', '--batch', batch.name, '--yes')
        assert code == 0 and result['changed'] == 1 and result['pushed'], result
        urls = csv_rows(batch / 'images.csv')
        assert len(urls) == 1 and urls[0]['slot'] == all_items[0]['slot']
        assert all(re.fullmatch(r'[a-z0-9/-]+\.jpg', r['url'].split('/main/')[1]) for r in urls)
        code, repeat = captured(ws, 'publish-images', '--batch', batch.name, '--yes')
        assert code == 0 and repeat['changed'] == 0 and not repeat['pushed']
        # New remote content is fast-forwarded, never removed or committed with our images.
        incoming = ws / 'incoming'
        local_git(ws, 'clone', bare.as_uri(), str(incoming))
        (incoming / 'other-batch.txt').write_text('别人发布的批次', 'utf-8')
        local_git(incoming, 'add', '.')
        local_git(incoming, 'commit', '-m', 'another batch')
        local_git(incoming, 'push', 'origin', 'main')
        first = all_items[0]
        source = output_dir(batch / 'images', first['group'], first['variante']) / (SLOTS.get(first['slot']) or 'variante.jpg')
        source.write_bytes(jpeg('green'))
        code, preview = captured(ws, 'publish-images', '--batch', batch.name)
        assert code == 2 and preview['count'] == 0, preview
        refreshed, _ = host.selection(batch, True)
        mark(batch, [refreshed[0]])
        code, result = captured(ws, 'publish-images', '--batch', batch.name, '--yes')
        assert code == 0 and result['changed'] == 1, result
        assert (clone / first['path']).read_bytes() == source.read_bytes()
        assert (clone / 'other-batch.txt').exists() and (clone / 'old-batch/legacy.jpg').read_bytes() == legacy
        assert urls[0]['url'] == csv_rows(batch / 'images.csv')[0]['url'], '同路径覆盖保持网址稳定'
        # Fail after committing, retain old CSV; retry safely publishes the pending commit.
        source.write_bytes(jpeg('yellow'))
        refreshed, _ = host.selection(batch, True)
        mark(batch, [refreshed[0]])
        previous_csv = (batch / 'images.csv').read_bytes()
        real_git = host.git
        def fail_push(directory, *args, **kwargs):
            if args and args[0] == 'push':
                raise Problem('网络不可用；请检查网络后重试。')
            return real_git(directory, *args, **kwargs)
        with patch.object(host, 'git', fail_push):
            code, result = captured(ws, 'publish-images', '--batch', batch.name, '--yes')
        assert code == 1 and (batch / 'images.csv').read_bytes() == previous_csv
        code, result = captured(ws, 'publish-images', '--batch', batch.name, '--yes')
        assert code == 0 and result['pushed'] and result['changed'] == 0, result
        # Interrupt before commit, then recover selected staged changes.
        source.write_bytes(jpeg('purple'))
        refreshed, _ = host.selection(batch, True)
        mark(batch, [refreshed[0]])
        def fail_commit(directory, *args, **kwargs):
            if args and args[0] == 'commit':
                raise Problem('提交中断；请重跑。')
            return real_git(directory, *args, **kwargs)
        with patch.object(host, 'git', fail_commit):
            code, result = captured(ws, 'publish-images', '--batch', batch.name, '--yes')
        assert code == 1 and (batch / 'images.csv').read_bytes() == previous_csv
        code, result = captured(ws, 'publish-images', '--batch', batch.name, '--yes')
        assert code == 0 and result['changed'] == 1 and result['pushed'], result
        # Unrelated dirty files and pending commits are never pushed.
        (clone / 'private.txt').write_text('unrelated', 'utf-8')
        code, result = captured(ws, 'publish-images', '--batch', batch.name, '--yes')
        assert code == 1 and any('无关' in e for e in result['errors'])
        (clone / 'private.txt').unlink()
        with patch.object(host, 'head_image', return_value=dict(url=urls[0]['url'],ok=False,status=404,content_type='text/html')):
            code, result = captured(ws, 'publish-images', '--batch', batch.name, '--yes')
        assert code == 1 and result['output'] and any('抽查失败' in e for e in result['errors']), result
        code, result = captured(ws, 'finalize', '--batch', batch.name)
        assert code == 0 and not result['upload_ready'] and not (ws / '台账.csv').exists(), result
        assert any('抽查' in w for w in result['warnings']) and any('占位' in w for w in result['warnings'])
        # Pending pictures may be included explicitly; failed groups never are.
        code, result = captured(ws, 'publish-images', '--batch', batch.name, '--include-unreviewed', '--yes')
        assert code == 0 and result['count'] == len(all_items) and len(result['checks']) == 5, result
        assert len(csv_rows(batch / 'images.csv')) == len(all_items)
        custom_shop = read_json(ws / 'shop.json')
        raw_base = custom_shop['image_host']['base_url']
        custom_shop['image_host']['base_url'] = 'https://cdn.example.invalid/images'
        json_write(ws / 'shop.json', custom_shop)
        code, result = captured(ws, 'publish-images', '--batch', batch.name, '--include-unreviewed', '--yes')
        assert code == 0 and not result['pushed'] and all(r['url'].startswith('https://cdn.example.invalid/images/') for r in csv_rows(batch / 'images.csv'))
        # Re-running setup preserves a user-selected CDN base_url.
        code, result = captured(ws, 'image-host-setup', '--repo', 'offline-owner/images', '--yes')
        assert code == 0 and read_json(ws / 'shop.json')['image_host']['base_url'] == 'https://cdn.example.invalid/images'
        custom_shop['image_host']['base_url'] = raw_base
        json_write(ws / 'shop.json', custom_shop)
        code, result = captured(ws, 'publish-images', '--batch', batch.name, '--include-unreviewed', '--yes')
        assert code == 0, result
        print('publish：仅已审字节、审阅失效、仅变化提交、同路径覆盖、远端共存、失败重试、跳过清单和 HEAD 抽查通过')
        assert host.slug(batch.name) == host.slug(batch.name) != host.slug(batch.name + '二')
        assert host.slug('ABC') != host.slug('abc') and host.slug('冬季') != host.slug('夏季')
        assert re.fullmatch('[a-z0-9-]+', host.slug(batch.name))
        # Block manufacturers without altering the previous draft output or ledger.
        good_shop = read_json(ws / 'shop.json')
        bad = read_json(ws / 'shop.json')
        bad['suppliers']['3321']['fabricante'] = '不存在的制造商'
        json_write(ws / 'shop.json', bad)
        old_output = (batch / 'output/miravia_upload.xlsm').read_bytes()
        code, result = captured(ws, 'finalize', '--batch', batch.name)
        assert code == 1 and any('fabricante' in e for e in result['errors'])
        assert (batch / 'output/miravia_upload.xlsm').read_bytes() == old_output and not (ws / '台账.csv').exists()
        json_write(ws / 'shop.json', good_shop)
        cli(ws, 'price', '--batch', batch.name)
        code, result = captured(ws, 'finalize', '--batch', batch.name, '--max-groups', '2')
        assert code == 0 and result['upload_ready'] and len(result['outputs']) == 2, result
        upload_outputs = result['outputs']
        notes = '自动决定：保留已审过的图。\n\n待办原文 | 不改格式。\n'
        (batch / 'notes.md').write_text(notes, 'utf-8')
        ledger_before = (ws / '台账.csv').read_bytes()
        code, report = captured(ws, 'report', '--batch', batch.name)
        assert code == 0 and report['upload_ready'] and report['outputs'] == upload_outputs, report
        assert report['count'] == 5 and report['groups'] == 3 and len(report['prices']) == 5
        assert (batch / '早上看这里.md').read_text('utf-8').endswith(notes)
        assert (ws / '台账.csv').read_bytes() == ledger_before
        template = ws / 'template' / good_shop['template']
        family_groups = []
        for output in result['outputs']:
            independent_zip(template, Path(output))
            wb = openpyxl.load_workbook(output, keep_vba=True)
            try:
                sheet = wb['Pantilla']
                groups = {sheet.cell(n,1).value for n in range(5,sheet.max_row+1)}
                assert len(groups) <= 2
                family_groups.append(groups)
            finally:
                wb.close()
        assert {'T01','T01B'} in family_groups and {'T02'} in family_groups, family_groups
        ledger = csv_rows(ws / '台账.csv')
        assert len(ledger) == 5 and all(r['批次'] == batch.name for r in ledger)
        assert ledger[0]['EAN'].startswith('2') and ledger[0]['卖家 SKU']
        assert ledger[0]['每单利润'] == csv_rows(batch / 'priced.csv')[0]['profit_with_coupon']
        # Keep another batch's ledger rows; overwrite matching EAN and preserve date.
        other = dict(ledger[0], **{'批次':'another batch', '日期':'2025-01-01'})
        ledger[0]['售价'] = '0'
        write_csv(ws / '台账.csv', [*ledger, other], list(ledger[0]))
        code, result = captured(ws, 'finalize', '--batch', batch.name, '--max-groups', '3')
        assert code == 0 and result['upload_ready'] and len(result['outputs']) == 1, result
        updated = csv_rows(ws / '台账.csv')
        assert len(updated) == 6 and other in updated
        assert sum(r['批次'] == batch.name and r['EAN'] == ledger[0]['EAN'] for r in updated) == 1
        assert not (batch / 'output/miravia_upload_2.xlsm').exists()
        assert next(r for r in updated if r['批次']==batch.name and r['EAN']==ledger[0]['EAN'])['售价'] != '0'
        code, result = captured(ws, 'finalize', '--batch', batch.name, '--max-groups', '1')
        assert code == 1 and any('不能拆开' in e for e in result['errors'])
        code, result = captured(ws, 'finalize', '--batch', batch.name, '--max-groups', '0')
        assert code == 1
        # Current source changes invalidate publication even though URLs still look valid.
        source.write_bytes(jpeg('black'))
        code, result = captured(ws, 'finalize', '--batch', batch.name)
        assert code == 0 and not result['upload_ready'] and any('成品已变化' in w for w in result['warnings'])
        code, report = captured(ws, 'report', '--batch', batch.name)
        assert code == 0 and not report['upload_ready'] and any('成品已变化' in t for t in report['todos']), report
        assert csv_rows(ws / '台账.csv') == updated
        # Rejected unrelated outgoing commits never reach the fake remote.
        remote_head = local_git(ws, '--git-dir', str(bare), 'rev-parse', 'main')
        (clone / 'private.txt').write_text('pending unrelated commit', 'utf-8')
        local_git(clone, 'add', '--', 'private.txt')
        local_git(clone, 'commit', '-m', 'unrelated local change')
        mark(batch, host.selection(batch, True)[0])
        code, result = captured(ws, 'publish-images', '--batch', batch.name, '--yes')
        assert code == 1 and any('待推送提交' in e for e in result['errors']), result
        code, result = captured(ws, 'image-host-setup', '--repo', 'offline-owner/images', '--yes')
        assert code == 1 and any('待推送提交' in e for e in result['errors']), result
        assert local_git(ws, '--git-dir', str(bare), 'rev-parse', 'main') == remote_head
        print('finalize：占位/抽查失败/制造商/旧发布阻断，完整及分表 ZIP 宏校验，主组+B 不拆，台账幂等覆盖与跨批次保留通过')
    # Thin HEAD boundary is tested without a socket: MIME and status both matter.
    class Response:
        status = 200
        def __init__(self, mime):
            from email.message import Message
            self.headers = Message()
            self.headers['Content-Type'] = mime
        def __enter__(self): return self
        def __exit__(self, *args): return False
    for mime, ready in [('image/jpeg',True),('text/html',False)]:
        with patch.object(host, 'urlopen', return_value=Response(mime)) as request:
            assert host.head_image('https://example.invalid/01-principal.jpg')['ok'] == ready
            assert request.call_args.args[0].method == 'HEAD'
    with patch.object(host, 'urlopen', side_effect=OSError('offline')):
        assert not host.head_image('https://example.invalid/01-principal.jpg')['ok']
    valid = dict(type='github', repo='offline-owner/images', branch='main', base_url='https://example.invalid/images')
    for field, bad in [('repo','../../repo'),('branch','../main'),('base_url','https://user:secret@example.invalid')]:
        try:
            host.validate_host(dict(valid, **{field:bad}))
            raise AssertionError('非法配置未阻止')
        except Problem:
            pass
    for cmd in ('image-host-setup', 'publish-images', 'finalize'):
        assert cli(ws, cmd, '--help')['help']
    print('PUBLISH OK')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    try:
        main()
    except Exception:
        import traceback
        traceback.print_exc()
        print('PUBLISH FAILED')
        raise SystemExit(1)
