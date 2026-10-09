"""A2 离线验收：只用假网页返回和本仓库模板副本，不登录、不联网。"""
import copy
import hashlib
import io
import json
import shutil
import sys
from contextlib import contextmanager, redirect_stdout
from unittest.mock import patch
from types import SimpleNamespace

import openpyxl
from PIL import Image

from helpers import FIXTURES, ROOT, cli, csv_rows, json_write, workspace
from listing_core.common import CANDIDATE_FIELDS, Problem, write_csv
from listing_core.workspace import load_shop
from listing_core.yollgo import fetch, read_barcodes, search
from listing_core.yollgo_browser import BrowserClient
from listing_core.yollgo_build import build


def code(prefix):
    """测试独立校验位：从右侧开始按 3、1 交替加权。"""
    return prefix + str((10 - sum(int(d) * (3 if i % 2 == 0 else 1)
                                  for i, d in enumerate(reversed(prefix))) % 10) % 10)


def fails(fn, *words):
    try:
        fn()
    except Problem as exc:
        assert all(word in str(exc) for word in words), (words, exc)
        return exc
    raise AssertionError('应当报错但成功了：' + str(words))


def setup(name):
    ws = workspace(name)
    cli(ws, 'init')
    shutil.copyfile(FIXTURES / 'shop.json', ws / 'shop.json')
    shop = json.loads((ws / 'shop.json').read_text('utf-8'))
    shutil.copyfile(FIXTURES / 'template' / shop['template'], ws / 'template' / shop['template'])
    cli(ws, 'new-batch', 'sample')
    return ws, ws / 'batches/sample', shop


def product(barcode, price=6.5, art='1000015560', minimum=1):
    return dict(artId=art, usercode=barcode, bianhao=barcode,
                namecn='MANTA BORREGO 130X160CM AZURITA', namees='',
                muluID='001067', baozhuangshu=minimum, precio=price,
                imageHash='f8e13d55325363057daa71f7c9317fee', is_attr=0, attributes=[],
                extra_public_field='保留未知原始商品字段')


class FakeClient:
    def __init__(self, quotes, shops=None):
        self.quotes = quotes
        self.queries, self.downloads = [], []
        self.public = shops or [dict(shopId=sid, name='商家 ' + sid, namees='EMPRESA ' + sid,
                                    des='公开说明', tel='公开电话', baseurl='https://img-eu-2.freex.es')
                                for sid in ('999', '3321', '027')]
        stream = io.BytesIO()
        Image.new('RGB', (600, 600), 'navy').save(stream, format='JPEG')
        self.image_bytes = stream.getvalue()

    def shops(self):
        return self.public

    def search(self, sid, keyword):
        self.queries.append((sid, keyword))
        return copy.deepcopy(self.quotes.get((sid, keyword), []))

    def image(self, url):
        self.downloads.append(url)
        assert '/600x600/' in url
        return self.image_bytes


def test_fetch():
    ws, batch, shop = setup('yollgo-fetch-')
    # Explicit order differs from followed-shop order. 027 is not followed:
    # it still must be queried first and win over a cheaper followed supplier.
    shop['suppliers'] = {sid: shop['suppliers'][sid] for sid in ('027', '3321')}
    main, unknown, zero, missing = [code('7654321000' + n) for n in ('10', '20', '30', '40')]
    (batch / 'barcodes.txt').write_text(f'# 批次注释\n\n{main}  原样  备注  \n{unknown}\n{zero}\n{missing}\n', encoding='utf-8')
    assert read_barcodes(batch / 'barcodes.txt')[0]['note'] == '原样  备注  '
    cli(ws, 'new-batch', 'older')
    write_csv(ws / 'batches/older/candidates.csv', [dict(group='G07', ean=code('765432199910'), unit_ean=main)], ['group', 'ean', 'unit_ean'])
    alias = product(main, 12.1, minimum=12)
    alias['bianhao'] = '其他货号'  # Exact usercode is also valid.
    quotes = {('027', main): [alias], ('3321', main): [product(main, 2)],
              ('999', main): [product(main, 1), product(main + '9', 0.1, art='false-match')],
              ('999', unknown): [product(unknown, 3, art='200')],
              ('999', zero): [product(zero, 0, art='300')]}
    quotes[('888', unknown)] = [product(unknown, 1, art='200')]
    followed = [s for s in FakeClient({}).public if s['shopId'] != '027']
    followed.append(dict(shopId='888', name='便宜店', namees='EMPRESA 888', baseurl='https://img-eu-2.freex.es'))
    client = FakeClient(quotes, followed)
    result = fetch(ws, batch, shop, client)
    data = json.loads((batch / 'yollgo.json').read_text('utf-8'))
    assert client.queries[:3] == [('027', main), ('3321', main), ('999', main)]
    assert data['items'][0]['selected'] == dict(shop_id='027', art_id='1000015560')
    assert len(data['items'][0]['offers']) == 3
    assert data['items'][0]['offers'][0]['product']['extra_public_field']
    assert data['items'][1]['selected']['shop_id'] == '888'
    assert result['issues']['not_found'] == [missing]
    assert result['issues']['zero_price'] == [zero]
    assert len(result['issues']['unregistered']) == 2
    assert result['issues']['bulk_minimum'][0]['quantity'] == 12
    assert result['issues']['already_listed'] == [dict(barcode=main, batch='older', group='G07')]
    assert all(any(word in warning for warning in result['warnings']) for word in ('制造商会空', '找不到', '价格为 0', '12', 'older'))
    assert (batch / 'src/3321-1000015560.jpg').read_bytes() == client.image_bytes
    assert (batch / 'src/999-1000015560.jpg').is_file(), '不同商家同 artId 必须分文件'
    assert (batch / 'src/_sheet.jpg').is_file()
    public = json.loads((ws / 'yollgo_shops.json').read_text('utf-8'))
    assert public['999']['namees'] == 'EMPRESA 999' and public['999']['tel'] == '公开电话'
    old_count = len(client.downloads)
    fetch(ws, batch, shop, client)
    assert len(client.downloads) == old_count, '已有原图不应重复下载'
    assert search(client, '999', unknown)['products'][0]['barcode'] == unknown
    # A no-match rerun must clear its stale overview.
    fetch(ws, batch, shop, FakeClient({}))
    assert not (batch / 'src/_sheet.jpg').exists()
    (batch / 'barcodes.txt').write_text('00001234 保留前导零\n00001234 重复', encoding='utf-8')
    fails(lambda: read_barcodes(batch / 'barcodes.txt'), '第 2 行', '重复')
    print('fetch：配置顺序、未关注配置店、最低价、多家报价、精确匹配、全部提醒、原图复用与公开资料通过')


def group_data(barcode, second=None):
    base = dict(barcode=barcode, weight_kg=1.2, len_cm=35, wid_cm=30, hei_cm=10, estimado=True)
    items = [dict(base, var1_value='130x160', var2_value='Crema'),
             dict(base, var1_value='130x160', var2_value='Marrón',
                  pack_measurements={'2': dict(weight_kg=2.8, len_cm=40, wid_cm=35, hei_cm=18, estimado=True)})]
    if second:
        items.append(dict(base, barcode=second, var1_value='160x220', var2_value='Crema'))
    return dict(version=1, groups=[dict(group='G01', var1_name='Tamaño', var2_name='Color',
                items=items, packs=[2], notes='按原图识别颜色')])


def test_build():
    ws, batch, shop = setup('yollgo-build-')
    barcode, second = code('765432100010'), code('765432100020')
    shop['suppliers']['3321']['price_includes_iva'] = True
    json_write(ws / 'shop.json', shop)
    fake = FakeClient({('3321', barcode): [product(barcode, 12.1, minimum=12)],
                       ('3321', second): [product(second, 6.05, art='1000015561')]})
    (batch / 'barcodes.txt').write_text(f'{barcode} 用户备注\n{second}', encoding='utf-8')
    fetch(ws, batch, shop, fake)
    data = group_data(barcode, second)
    json_write(batch / 'groups.json', data)
    cli(ws, 'new-batch', 'other')
    collision = code(barcode[:11] + '1')
    write_csv(ws / 'batches/other/candidates.csv', [dict(group='G09', ean=collision)], ['group', 'ean'])
    template = ws / 'template' / shop['template']
    before = hashlib.sha256(template.read_bytes()).hexdigest()
    result = cli(ws, 'build', '--batch', 'sample')
    assert result['count'] == 6 and result['groups'] == 2
    rows, maps = csv_rows(batch / 'candidates.csv'), csv_rows(batch / 'variantes.csv')
    assert list(rows[0]) == CANDIDATE_FIELDS
    assert rows[0]['ean'] == barcode and rows[1]['ean'] == code(barcode[:11] + '2')
    assert rows[1]['unit_ean'] == barcode and maps[1]['variante'] == '130x160-marron'
    assert all(float(row['unit_cost_ex_iva']) == 10 for row in rows[:2])
    assert float(rows[3]['unit_cost_ex_iva']) == 20 and rows[3]['pack_qty'] == '2'
    assert float(rows[3]['weight_kg']) == 2.4 and float(rows[3]['hei_cm']) == 20
    assert float(rows[4]['weight_kg']) == 2.8 and float(rows[4]['hei_cm']) == 18
    assert rows[4]['ean'] == code('22' + rows[1]['ean'][-10:])
    assert all(row['unit_ean'] in (barcode, second) for row in rows)
    assert len({r['ean'] for r in rows}) == 6
    assert '用户备注' in rows[0]['notes'] and '12' in rows[0]['notes'] and '估计' in rows[0]['notes']
    assert all(row['market_low'] == row['market_high'] == '' for row in rows)
    assert hashlib.sha256(template.read_bytes()).hexdigest() == before
    cli(ws, 'price', '--batch', 'sample')
    priced = csv_rows(batch / 'priced.csv')
    assert float(priced[0]['real_cost']) == 12.62
    assert float(priced[3]['real_cost']) == 25.24, '组合装成本只能乘一次件数'
    assert all(row['vs_market'] == '' for row in priced)
    snapshot = {name: (batch / name).read_bytes() for name in ('candidates.csv', 'variantes.csv', 'build.json')}
    rejected = cli(ws, 'build', '--batch', 'sample', ok=False)
    assert '--force' in rejected['errors'][0]
    assert all((batch / name).read_bytes() == value for name, value in snapshot.items())
    # Both independently invalid items should be reported with group/item paths.
    bad = copy.deepcopy(data)
    del bad['groups'][0]['items'][0]['weight_kg']
    bad['groups'][0]['items'][1]['estimado'] = 'true'
    json_write(batch / 'groups.json', bad)
    rejected = cli(ws, 'build', '--batch', 'sample', '--force', ok=False)
    assert any('G01' in e and '第 1 项' in e and 'weight_kg' in e for e in rejected['errors'])
    assert any('G01' in e and '第 2 项' in e and 'estimado' in e for e in rejected['errors'])
    assert all((batch / name).read_bytes() == value for name, value in snapshot.items())
    json_write(batch / 'groups.json', data)
    raw_snapshot = (batch / 'yollgo.json').read_bytes()
    invalid_quotes = json.loads(raw_snapshot)
    invalid_quotes['items'][0]['selected'] = None
    json_write(batch / 'yollgo.json', invalid_quotes)
    fails(lambda: build(ws, batch, shop, force=True), 'G01', '第 1 项', '选中的报价')
    (batch / 'yollgo.json').write_bytes(raw_snapshot)
    # Exhaust every color id in other batches.
    write_csv(ws / 'batches/other/candidates.csv', [dict(group='G09', ean=code(barcode[:11] + str(i))) for i in range(1, 10)], ['group', 'ean'])
    fails(lambda: build(ws, batch, shop, force=True), 'G01', '1～9', '条码')
    pack_code = code('22' + barcode[-10:])
    write_csv(ws / 'batches/other/candidates.csv', [dict(group='G09', ean=pack_code)], ['group', 'ean'])
    fails(lambda: build(ws, batch, shop, force=True), 'G01B', pack_code, '撞号')
    write_csv(ws / 'batches/other/candidates.csv', [dict(group='G09', ean=barcode)], ['group', 'ean'])
    fails(lambda: build(ws, batch, shop, force=True), barcode, 'other', 'G09')
    # Template collision sources: hidden validation list and visible Pantilla.
    write_csv(ws / 'batches/other/candidates.csv', [], ['group', 'ean'])
    backup = template.read_bytes()
    wb = openpyxl.load_workbook(template, keep_vba=True)
    from listing_core.headers import hidden_column
    col = hidden_column(wb, 'ean', 'global_validation_hide', 1)
    wb['global_validation_hide'].cell(3, col).value = barcode
    wb.save(template); wb.close()
    fails(lambda: build(ws, batch, shop, force=True), barcode, '模板')
    template.write_bytes(backup)
    # A later source EAN with the same first 11 digits takes precedence over
    # generated colors, regardless of item order in groups.json.
    later = code(barcode[:11] + '2')
    raw_before = (batch / 'yollgo.json').read_bytes()
    raw = json.loads(raw_before)
    later_item = copy.deepcopy(raw['items'][1])
    later_item['barcode'] = later
    later_item['offers'][0]['product']['bianhao'] = later
    later_item['offers'][0]['product']['usercode'] = later
    raw['items'].append(later_item)
    json_write(batch / 'yollgo.json', raw)
    json_write(batch / 'groups.json', group_data(barcode, later))
    build(ws, batch, shop, force=True)
    current = csv_rows(batch / 'candidates.csv')
    assert current[1]['ean'] == code(barcode[:11] + '1') and current[2]['ean'] == later
    # Two distinct originals produce identical pack suffixes: fail before any
    # output mutation, with the conflicting pack barcode in the message.
    collision_item = copy.deepcopy(later_item)
    collision_source = code('925' + barcode[3:12])
    assert collision_source != barcode and collision_source[-10:] == barcode[-10:]
    collision_item['barcode'] = collision_source
    collision_item['selected']['art_id'] = 'other-art'
    collision_item['offers'][0]['product']['artId'] = 'other-art'
    raw['items'].append(collision_item)
    json_write(batch / 'yollgo.json', raw)
    json_write(batch / 'groups.json', group_data(barcode, collision_source))
    fails(lambda: build(ws, batch, shop, force=True), 'G01B', '撞号')
    (batch / 'yollgo.json').write_bytes(raw_before)
    json_write(batch / 'groups.json', data)
    for name, value in snapshot.items(): (batch / name).write_bytes(value)
    wb = openpyxl.load_workbook(template, keep_vba=True)
    sheet = wb['Pantilla']; col = next(c.column for c in sheet[1] if c.value == 'Código EAN')
    sheet.cell(5, col).value = barcode; wb.save(template); wb.close()
    fails(lambda: build(ws, batch, shop, force=True), barcode, 'Pantilla')
    template.write_bytes(backup)
    # Guard against interrupted/manual edits changing the cost interpretation.
    (batch / 'candidates.csv').write_bytes(snapshot['candidates.csv'] + b'\n')
    result = cli(ws, 'price', '--batch', 'sample', ok=False)
    assert 'build.json' in result['errors'][0]
    (batch / 'candidates.csv').write_bytes(snapshot['candidates.csv'])
    assert load_shop(ws)['suppliers']['3321']['price_includes_iva'] is True
    bad_shop = copy.deepcopy(shop); bad_shop['suppliers']['3321']['price_includes_iva'] = 'true'
    json_write(ws / 'shop.json', bad_shop)
    fails(lambda: load_shop(ws), 'suppliers.3321.price_includes_iva')
    json_write(ws / 'shop.json', shop)
    print('build：颜色与组合装条码、跨批次及模板撞号、颜色号耗尽、IVA、重量尺寸、覆盖保护和分组报错通过')
    print('price：市场价留空、整包成本只乘一次、生成中断校验通过')


def test_first_build_interruption():
    ws, batch, shop = setup('yollgo-interrupt-')
    barcode = code('765432100010')
    (batch / 'barcodes.txt').write_text(barcode, encoding='utf-8')
    fetch(ws, batch, shop, FakeClient({('3321', barcode): [product(barcode)]}))
    json_write(batch / 'groups.json', group_data(barcode))
    original = (batch / 'candidates.csv').read_bytes()
    with patch('listing_core.yollgo_build.write_csv', side_effect=OSError('离线模拟写入中断')):
        try:
            build(ws, batch, shop)
        except OSError:
            pass
        else:
            raise AssertionError('中断模拟未生效')
    assert (batch / 'candidates.csv').read_bytes() == original
    result = cli(ws, 'price', '--batch', 'sample', ok=False)
    assert 'build.json' in result['errors'][0]
    cli(ws, 'build', '--batch', 'sample')
    cli(ws, 'price', '--batch', 'sample')
    print('中断恢复：首次生成中断不会误算组合装成本，可直接重跑 build 恢复')


def test_cli_and_browser():
    ws, batch, shop = setup('yollgo-cli-')
    for command in ('yollgo-login', 'fetch', 'yollgo-search', 'build'):
        assert cli(ws, command, '--help')['help']
    barcode = code('765432100010')
    fake = FakeClient({('3321', barcode): [product(barcode)]})
    (batch / 'barcodes.txt').write_text(barcode, encoding='utf-8')
    import mlist
    @contextmanager
    def fake_session(*args, **kwargs):
        yield fake
    with patch('listing_core.yollgo_browser.session', fake_session):
        stream = io.StringIO()
        with redirect_stdout(stream):
            assert mlist.main(['fetch', '--ws', str(ws), '--batch', 'sample', '--json']) == 0
        assert json.loads(stream.getvalue())['found'] == 1
        stream = io.StringIO()
        with redirect_stdout(stream):
            assert mlist.main(['yollgo-search', '--ws', str(ws), '--shop', '3321', barcode]) == 0
        assert barcode in stream.getvalue() and '€6.5' in stream.getvalue()
    class FakePage:
        def __init__(self): self.calls = []
        def evaluate(self, script, params):
            self.calls.append(params)
            offset = params['params']['from']
            return dict(ok=True, items=[product(str(i), art=str(i)) for i in range(20)] if offset == 0 else [product(barcode)])
    page = FakePage()
    assert len(BrowserClient(page).search('027', 'manta')) == 21
    assert [p['params']['from'] for p in page.calls] == [0, 20]
    class FailedPage:
        def evaluate(self, *args): return dict(ok=False)
    fails(lambda: BrowserClient(FailedPage()).shops(), '友购', 'yollgo-login')
    # Replace Playwright itself: verify persistent Edge, login detection, window
    # close and visible/headless choices without launching any real browser.
    from listing_core.yollgo_browser import session, login
    calls, messages = [], []
    # 友购关窗后登录失效(10-09 实测):一律开可见窗口,没登录就在同一窗口里等用户登录。
    user = {'id': -1, 'alive': False, 'closed': False}
    class SessionPage:
        def goto(self, url, **kwargs): assert url == 'https://app.yollgo.com'
        def wait_for_function(self, script, **kwargs): pass
        def wait_for_timeout(self, ms): pass
        def is_closed(self): return user['closed']
        def evaluate(self, script):
            return user['alive'] if 'shops()' in script else user['id']
    class FakeContext:
        def __init__(self): self.pages = [SessionPage()]
        def close(self): calls.append('closed')
    class FakePlaywright:
        def __enter__(self):
            return SimpleNamespace(chromium=SimpleNamespace(launch_persistent_context=self.launch))
        def __exit__(self, *args): pass
        def launch(self, directory, **kwargs):
            assert directory == str(ws / '.yollgo-browser')
            calls.append(kwargs)
            return FakeContext()
    module = SimpleNamespace(sync_playwright=FakePlaywright)
    with patch.dict(sys.modules, {'playwright.sync_api': module}):
        user['closed'] = True  # 用户没登录就关了窗口
        def closed_before_login():
            with session(ws, notify=messages.append):
                raise AssertionError('未登录时不应进入业务层')
        fails(closed_before_login, '窗口已关闭', '重新运行')
        assert calls == [dict(channel='msedge', headless=False), 'closed']
        assert messages and '登录' in messages[0]
        calls.clear(); messages.clear()
        user.update(id='1000436996', alive=True, closed=False)
        with session(ws, notify=messages.append) as client:
            assert client.page is not None
        assert messages == [] and calls == [dict(channel='msedge', headless=False), 'closed']
        calls.clear()
        assert '已登录' in login(ws, notify=messages.append)['message']
        assert calls == [dict(channel='msedge', headless=False), 'closed']
    import builtins
    actual_import = builtins.__import__
    def without_playwright(name, *args, **kwargs):
        if name == 'playwright.sync_api': raise ImportError('offline missing dependency')
        return actual_import(name, *args, **kwargs)
    with patch('builtins.__import__', without_playwright):
        fails(lambda: login(ws), 'pip install', '不要运行 playwright install')
        assert cli(ws, 'build', '--help')['help']  # Non-browser commands stay usable.
    print('接口：假浏览器 CLI、搜索输出、分页、查询失败和四个新命令帮助通过（无登录、无联网）')


def main():
    test_fetch()
    test_build()
    test_first_build_interruption()
    test_cli_and_browser()
    print('YOLLGO OK')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
