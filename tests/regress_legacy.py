"""只读旧基准；新流程和所有临时文件均在本仓库 .tmp/。"""
import json
import sys
from pathlib import Path
from decimal import Decimal, InvalidOperation
import re
import openpyxl
from helpers import ROOT, cli, csv_rows, workspace

LEGACY = Path(r'E:\我的文件\桌面\Miravia')


IMG = re.compile(r'<p><img [^>]*/></p>')  # 描述里自动插的图床图,旧基准没有


def strip_img(v):
    return IMG.sub('', v) if isinstance(v, str) else v


def equal_cell(a, b):
    if a == b: return True
    try: return Decimal(a) == Decimal(b)
    except (InvalidOperation, ValueError, TypeError): return False


def main():
    if not LEGACY.is_dir():
        print('SKIP：旧项目不存在。')
        return 0
    for rel in ('data/priced.csv','output/miravia_upload.xlsm'):
        assert (LEGACY/rel).is_file(), f'旧基准缺失：{rel}；请提供基准，不运行或修改旧项目。'
    ws = workspace('regress-')
    cli(ws,'import-legacy','--from',LEGACY,'--batch','legacy')
    result=cli(ws,'price','--batch','legacy')
    print(f"price：{result['count']} 行；{len(result['warnings'])} 条提醒")
    # Only enrich the temporary imported copy with synthetic image-test facts.
    # They never enter workbook cells or alter the read-only legacy baseline.
    from helpers import FIXTURES, json_write
    content_path = ws / 'batches/legacy/content.json'
    content = json.loads(content_path.read_text('utf-8-sig'))
    fixture = json.loads((FIXTURES / 'batch/content.json').read_text('utf-8-sig'))['T01']
    for item in content.values():
        for key in ('use_scene', 'hero_feature', 'scene_briefs'):
            item[key] = fixture[key]
    json_write(content_path, content)
    result=cli(ws,'fill','--batch','legacy')
    print(f"fill：{result['count']} SKU / {result['groups']} 组；{len(result['warnings'])} 条提醒")
    batch=ws/'batches/legacy'
    old,new=csv_rows(LEGACY/'data/priced.csv'),csv_rows(batch/'priced.csv')
    diffs=[]
    if len(old)!=len(new):diffs.append(f'priced.csv 行数：旧={len(old)} 新={len(new)}')
    for n,(a,b) in enumerate(zip(old,new),2):
        if list(a)!=list(b):diffs.append(f'priced.csv 第 {n} 行表头/列序不同')
        for key in sorted(set(a)|set(b)):
            if not equal_cell(a.get(key),b.get(key)):diffs.append(f'priced.csv 第 {n} 行 {key}：旧={a.get(key)!r} 新={b.get(key)!r}')
    wb1=openpyxl.load_workbook(LEGACY/'output/miravia_upload.xlsm',keep_vba=True)
    wb2=openpyxl.load_workbook(batch/'output/miravia_upload.xlsm',keep_vba=True)
    try:
        a,b=wb1['Pantilla'],wb2['Pantilla']
        if (a.max_row,a.max_column)!=(b.max_row,b.max_column):diffs.append('Pantilla 行列数不同')
        for r in range(1,max(a.max_row,b.max_row)+1):
            for c in range(1,max(a.max_column,b.max_column)+1):
                if a.cell(r,c).value!=strip_img(b.cell(r,c).value):
                    diffs.append(f'Pantilla {a.cell(r,c).coordinate}：旧={a.cell(r,c).value!r} 新={b.cell(r,c).value!r}')
    finally:
        wb1.close();wb2.close()
    if diffs:
        for d in diffs:print(d)
        print(f'REGRESS FAILED：{len(diffs)} 处差异；新工作区：{ws}')
        return 1
    print(f'priced.csv：{len(old)} 行 × {len(old[0])} 列，全部一致')
    print(f'Pantilla：{a.max_row} 行 × {a.max_column} 列，每个单元格值全部一致')
    print('REGRESS OK')
    # Deliberately retain temp workspace: never recursively delete its junction.
    return 0


if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    try: raise SystemExit(main())
    except Exception as exc:
        print(f'REGRESS FAILED：{exc}')
        raise SystemExit(1)
